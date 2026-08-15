"""Spreadsheet tool: CSV (stdlib) + XLSX (openpyxl extra) (ROADMAP 2027 H1)."""
from __future__ import annotations

import json

import pytest
from maverick.tools import spreadsheet as ss

# ---- CSV (stdlib, always available) -----------------------------------------

def test_csv_write_read_info(tmp_path):
    p = tmp_path / "data.csv"
    assert ss._run({"op": "write", "path": str(p),
                    "rows": [["a", "b"], [1, 2], [3, 4]]}).startswith("wrote 3")
    rows = json.loads(ss._run({"op": "read", "path": str(p)}))
    assert rows[0] == ["a", "b"] and rows[1] == ["1", "2"]   # csv reads back as strings
    assert "rows=3" in ss._run({"op": "info", "path": str(p)})


def test_neutralize_formula_helper():
    assert ss._neutralize_formula("=1+1") == "'=1+1"
    assert ss._neutralize_formula("+x") == "'+x"
    assert ss._neutralize_formula("-x") == "'-x"
    assert ss._neutralize_formula("@x") == "'@x"
    assert ss._neutralize_formula("\tx") == "'\tx"
    assert ss._neutralize_formula("\rx") == "'\rx"
    assert ss._neutralize_formula("\nx") == "'\nx"
    assert ss._neutralize_formula("  =x") == "'  =x"
    assert ss._neutralize_formula(" \t=x") == "' \t=x"
    assert ss._neutralize_formula("safe") == "safe"
    assert ss._neutralize_formula(42) == 42       # numbers untouched
    assert ss._neutralize_formula("") == ""       # empty untouched


def test_csv_write_neutralizes_formula_injection(tmp_path):
    p = tmp_path / "evil.csv"
    rows = [["=HYPERLINK(1)", "+1", "-2", "@SUM(A1)", "safe"]]
    ss._run({"op": "write", "path": str(p), "rows": rows})
    out = json.loads(ss._run({"op": "read", "path": str(p)}))
    # Every formula-leading cell is prefixed with ' so a spreadsheet app renders
    # it as text; benign cells are unchanged.
    assert out[0] == ["'=HYPERLINK(1)", "'+1", "'-2", "'@SUM(A1)", "safe"]


def test_set_cell_rejects_csv(tmp_path):
    p = tmp_path / "data.csv"
    ss._run({"op": "write", "path": str(p), "rows": [["x"]]})
    assert "xlsx-only" in ss._run({"op": "set_cell", "path": str(p), "cell": "A1", "value": 1})


def test_errors():
    assert ss._run({"op": "read"}).startswith("ERROR: path is required")
    assert ss._run({"op": "", "path": "x"}).startswith("ERROR: op is required")
    assert ss._run({"op": "read", "path": "/nope/missing.csv"}).startswith("ERROR: no such file")
    assert ss._run({"op": "bogus", "path": "x.csv"}).startswith("ERROR")


def test_confines_paths_to_sandbox_workspace(tmp_path):
    class _SB:
        workdir = tmp_path / "ws"

    (tmp_path / "ws").mkdir()
    secret = tmp_path / "secret.csv"
    secret.write_text("token,value\nHOST_SECRET,42\n")

    tool = ss.spreadsheet(_SB())
    assert "escapes the workspace" in tool.fn({"op": "read", "path": str(secret)})
    assert "escapes the workspace" in tool.fn({
        "op": "write", "path": "../out.csv", "rows": [["x"]],
    })
    assert not (tmp_path / "out.csv").exists()

    out = tool.fn({
        "op": "write", "path": "nested/data.csv", "rows": [["a", "b"], [1, 2]],
    })
    assert out.startswith("wrote 2")
    assert (tmp_path / "ws" / "nested" / "data.csv").exists()
    rows = json.loads(tool.fn({"op": "read", "path": "nested/data.csv"}))
    assert rows == [["a", "b"], ["1", "2"]]


# ---- XLSX (needs openpyxl) --------------------------------------------------

def test_xlsx_write_read_setcell_info(tmp_path):
    pytest.importorskip("openpyxl")
    p = tmp_path / "book.xlsx"
    assert ss._run({"op": "write", "path": str(p),
                    "rows": [["h1", "h2"], [10, 20]], "sheet": "Data"}).startswith("wrote 2")
    rows = json.loads(ss._run({"op": "read", "path": str(p), "sheet": "Data"}))
    assert rows[0] == ["h1", "h2"] and rows[1] == [10, 20]   # xlsx preserves int types
    assert ss._run({"op": "set_cell", "path": str(p), "cell": "A3",
                    "value": 99, "sheet": "Data"}).startswith("set A3")
    rows2 = json.loads(ss._run({"op": "read", "path": str(p), "sheet": "Data"}))
    assert rows2[2][0] == 99
    assert "xlsx sheets" in ss._run({"op": "info", "path": str(p)})


def test_set_cell_creates_missing_sheet_instead_of_renaming(tmp_path):
    # Regression: set_cell on an existing workbook used to RENAME the active
    # sheet to the requested name (silently corrupting the original sheet and
    # breaking cross-sheet formula references) instead of creating it.
    pytest.importorskip("openpyxl")
    import openpyxl
    p = tmp_path / "book.xlsx"
    ss._run({"op": "write", "path": str(p),
             "rows": [["h1"], [1]], "sheet": "Data"})
    out = ss._run({"op": "set_cell", "path": str(p), "cell": "B2",
                   "value": 42, "sheet": "Summary"})
    assert out.startswith("set B2")
    wb = openpyxl.load_workbook(p)
    assert wb.sheetnames == ["Data", "Summary"]   # Data survives, Summary created
    assert wb["Summary"]["B2"].value == 42
    assert wb["Data"]["A2"].value == 1            # original data untouched
    assert wb["Data"]["B2"].value is None         # not written into Data's grid


def test_set_cell_names_sheet_on_fresh_workbook(tmp_path):
    # Brand-new workbook: the default "Sheet" is still renamed to the request.
    pytest.importorskip("openpyxl")
    import openpyxl
    p = tmp_path / "new.xlsx"
    assert ss._run({"op": "set_cell", "path": str(p), "cell": "A1",
                    "value": 7, "sheet": "Only"}).startswith("set A1")
    wb = openpyxl.load_workbook(p)
    assert wb.sheetnames == ["Only"]
    assert wb["Only"]["A1"].value == 7


def test_xlsx_graceful_without_openpyxl(tmp_path, monkeypatch):
    # If openpyxl is unavailable, xlsx ops fail with an actionable install hint
    # rather than a raw ImportError.
    def _boom():
        raise RuntimeError("openpyxl not installed; .xlsx support needs it. "
                           "Run: pip install 'maverick-agent[spreadsheet]'")
    monkeypatch.setattr(ss, "_openpyxl", _boom)
    out = ss._run({"op": "write", "path": str(tmp_path / "x.xlsx"), "rows": [["a"]]})
    assert out.startswith("ERROR") and "maverick-agent[spreadsheet]" in out


def test_factory_shape():
    t = ss.spreadsheet()
    assert t.name == "spreadsheet" and callable(t.fn)
